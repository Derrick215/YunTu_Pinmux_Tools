#!/usr/bin/env python3
"""
Convert YTM32 PINMUX Excel files (.xlsx) to the same JSON format as the MD parser.
Reads IOMUX sheet from Excel files and outputs pin data for pinmux_data.js.
Usage: python convert_xlsx_to_json.py
"""
import os
import re
import json
from pathlib import Path
import openpyxl

XL_DIR = Path(__file__).parent / "YTM32_PINMUX_MD_Files"


def classify_pin_type(name):
    """Classify pin type based on its name."""
    if not name:
        return "UNKNOWN"
    name_upper = name.strip()
    if re.match(r'^(VDD\d*|VDDA|VREF[HRL]|VDD\d+)$', name_upper):
        return "POWER"
    if re.match(r'^(VSS\d*|VSSA|VREFL)$', name_upper):
        return "GND"
    if re.match(r'^(EXTAL\d*|XTAL\d*)$', name_upper, re.IGNORECASE):
        return "CLOCK"
    if re.match(r'^(JTAG_|RCU_RESET|SWD_|SWO)', name_upper, re.IGNORECASE):
        return "DEBUG"
    if re.match(r'^PT[A-Z]_', name_upper):
        return "GPIO"
    return "OTHER"


def parse_xlsx_file(filepath):
    """Parse a PINMUX Excel file and return chip data dict."""
    wb = openpyxl.load_workbook(filepath)

    # Find the IOMUX sheet (may be named 'IOMUX' or contain 'IOMUX')
    iomux_sheet = None
    for name in wb.sheetnames:
        if 'IOMUX' in name.upper():
            iomux_sheet = name
            break
    if not iomux_sheet:
        print(f"  WARNING: No IOMUX sheet found in {filepath.name}")
        return None

    ws = wb[iomux_sheet]

    # Extract chip name from filename: YTM32B1MC0x_PINMUX_V1.1.xlsx -> YTM32B1MC0x
    stem = filepath.stem
    chip_name = stem.split('_')[0]

    # Get version from Version sheet or filename
    version = None
    if 'Version' in wb.sheetnames:
        vs = wb['Version']
        for row in vs.iter_rows(min_row=2, max_row=vs.max_row, values_only=True):
            if row[0]:
                ver_str = str(row[0]).strip()
                ver_match = re.match(r'v?(\d+\.\d+)', ver_str, re.IGNORECASE)
                if ver_match:
                    version = ver_match.group(1)  # Take the last (latest) version
    if not version:
        version = stem.split('_V')[-1] if '_V' in stem else '1.0'

    # Parse header row to identify columns
    header_row = None
    for row in ws.iter_rows(min_row=1, max_row=5, values_only=False):
        # Check if this looks like a header row (contains 'NAME' or 'ALT')
        row_values = [str(c.value).strip() if c.value else '' for c in row]
        has_name = any('NAME' in v.upper() for v in row_values)
        has_alt = any(re.match(r'^ALT\d*$', v.upper()) for v in row_values)
        has_package = any(re.search(r'(LQFP|QFN|BGA|WLCSP)', v.upper()) for v in row_values)
        if has_name and (has_alt or has_package):
            header_row = row
            break

    if not header_row:
        print(f"  WARNING: No valid header row found in {filepath.name}")
        return None

    headers = [str(c.value).strip() if c.value else '' for c in header_row]

    # Identify column roles
    package_cols = {}  # {col_index: package_name}
    name_idx = None
    alt_indices = []
    drive_idx = None
    pull_up_idx = None
    slew_rate_idx = None
    pin_filter_idx = None

    for idx, h in enumerate(headers):
        h_upper = h.upper()
        if re.match(r'^ALT\d+$', h_upper):
            alt_indices.append(idx)
        elif re.search(r'(LQFP|QFN|BGA|WLCSP)', h_upper):
            package_cols[idx] = h
        elif h_upper == 'NAME':
            name_idx = idx
        elif 'HIGH DRIVE' in h_upper or 'SUPPORT HIGH DRIVE' in h_upper:
            drive_idx = idx
        elif 'PULLUP' in h_upper or 'PULLDOWN' in h_upper:
            pull_up_idx = idx
        elif 'SLEW' in h_upper:
            slew_rate_idx = idx
        elif 'FILTER' in h_upper:
            pin_filter_idx = idx

    # Also check for package columns that might not have been caught (e.g., just numbers)
    if not package_cols:
        for idx, h in enumerate(headers):
            if re.match(r'^\d+', h) or 'PIN' in h.upper():
                package_cols[idx] = h

    # Find the data start row (first row after header with a valid pin name)
    data_start = None
    for row_idx in range(header_row[0].row, ws.max_row + 1):
        row = list(ws.iter_rows(min_row=row_idx, max_row=row_idx, values_only=True))[0]
        if not row or name_idx is None or name_idx >= len(row):
            continue
        pin_name = str(row[name_idx]).strip() if row[name_idx] else ''
        # Skip header row
        if pin_name.upper() == 'NAME':
            continue
        if pin_name and re.match(r'^(PT[A-Z]_|VDD|VSS|VDDA|VREF|EXTAL|XTAL|JTAG_|RCU_|SWD_|SWO|NC)', pin_name):
            data_start = row_idx
            break

    if data_start is None:
        data_start = header_row[0].row + 1  # fallback

    # Parse data rows
    pins = []
    for row in ws.iter_rows(min_row=data_start, max_row=ws.max_row, values_only=True):
        if not row or name_idx is None or name_idx >= len(row):
            continue

        pin_name = str(row[name_idx]).strip() if row[name_idx] else ''
        if not pin_name:
            continue

        # Collect ALT functions (sorted by ALT index)
        alt_functions = []
        for ai in sorted(alt_indices):
            if ai < len(row) and row[ai] and str(row[ai]).strip():
                alt_val = str(row[ai]).strip()
                alt_val = re.sub(r'\s*WKU\[\d+\]\s*', '', alt_val).strip()
                if alt_val and alt_val not in alt_functions:
                    alt_functions.append(alt_val)

        # Collect package pin numbers
        package_pins = {}
        for pi, pkg_name in sorted(package_cols.items()):
            if pi < len(row) and row[pi] is not None and str(row[pi]).strip():
                val = str(row[pi]).strip()
                try:
                    package_pins[pkg_name] = int(val)
                except ValueError:
                    package_pins[pkg_name] = val

        # Determine pin type
        pin_type = classify_pin_type(pin_name)

        pin_obj = {
            "name": pin_name,
            "type": pin_type,
            "packagePins": package_pins,
            "altFunctions": alt_functions,
        }

        # Drive (high drive support)
        drive = None
        if drive_idx is not None and drive_idx < len(row) and row[drive_idx]:
            drive_val = str(row[drive_idx]).strip().upper()
            if drive_val == 'Y':
                drive = 'NORMAL'

        # Default state for NC pins
        default_state = None
        if pin_name == 'NC':
            default_state = 'NC'

        if drive:
            pin_obj["drive"] = drive
        if default_state:
            pin_obj["defaultState"] = default_state

        pins.append(pin_obj)

    # Collect all package names
    all_packages = set()
    for p in pins:
        all_packages.update(p["packagePins"].keys())

    sorted_packages = sorted(all_packages, key=lambda x: (
        -max((p["packagePins"].get(x, 0) for p in pins if x in p["packagePins"] and isinstance(p["packagePins"].get(x), int)), default=0)
    ))

    return {
        "name": chip_name,
        "version": version,
        "packages": sorted_packages,
        "pins": pins,
    }


def main():
    xl_files = sorted(XL_DIR.glob("*.xlsx"))
    if not xl_files:
        print("No Excel files found.")
        return

    chips = {}
    for filepath in xl_files:
        print(f"Parsing: {filepath.name} ... ", end="", flush=True)
        data = parse_xlsx_file(filepath)
        if data:
            chips[data["name"]] = data
            print(f"OK ({len(data['pins'])} pins, packages: {data['packages']})")
        else:
            print("FAILED")

    if chips:
        output_path = Path(__file__).parent / "pinmux_data_xlsx.js"
        json_str = json.dumps(chips, indent=2, ensure_ascii=False)
        js_content = f"// Auto-generated from Excel files\n// Generated by convert_xlsx_to_json.py\nconst PINMUX_DATA_XLSX = {json_str};\n"
        with open(output_path, 'w', encoding='utf-8') as f:
            f.write(js_content)
        print(f"\nOutput written to: {output_path}")
        print(f"Chips: {list(chips.keys())}")
    else:
        print("No data extracted.")


if __name__ == "__main__":
    main()
